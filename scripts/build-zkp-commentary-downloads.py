from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path("guides/zkp-commentary-usage/downloads")
OUT.mkdir(parents=True, exist_ok=True)
NAVY="0C2038"; GOLD="B9985B"; MUTED="58687A"

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcPr.append(shd)

def borders(cell, color="D9E1E8"):
    tcPr=cell._tc.get_or_add_tcPr()
    tcBorders=tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders=OxmlElement("w:tcBorders"); tcPr.append(tcBorders)
    for edge in ("top","left","bottom","right"):
        tag="w:"+edge; el=tcBorders.find(qn(tag))
        if el is None: el=OxmlElement(tag); tcBorders.append(el)
        el.set(qn("w:val"),"single"); el.set(qn("w:sz"),"4"); el.set(qn("w:color"),color)

def base_doc(title, subtitle="", author="Адвокат Зоран Стојанкиќ", version="2026"):
    d=Document(); sec=d.sections[0]
    sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
    normal=d.styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.15
    for name,size,color in [("Title",25,NAVY),("Heading 1",17,NAVY),("Heading 2",13,NAVY)]:
        s=d.styles[name]; s.font.name="Aptos"; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("AI ADVOKAT"); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(GOLD)
    p=d.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(title)
    if subtitle:
        p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(subtitle); r.italic=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(MUTED)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(f"{author} · {version}"); r.bold=True; r.font.size=Pt(10)
    d.add_paragraph("")
    return d

def notice(d,text,fill="FFF8E9",label=""):
    t=d.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    c=t.cell(0,0); shade(c,fill); borders(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=c.paragraphs[0]
    if label:
        r=p.add_run(label+" "); r.bold=True
    p.add_run(text); d.add_paragraph("")

def bullets(d,items,numbered=False):
    style="List Number" if numbered else "List Bullet"
    for item in items: d.add_paragraph(item,style=style)

def table(d,headers,rows):
    t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=True
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,NAVY); borders(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r=c.paragraphs[0].add_run(h); r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            c=cells[i]; borders(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; c.paragraphs[0].add_run(str(v))
    d.add_paragraph("")

def save(d,name):
    d.save(OUT/name)

d=base_doc("Водич за користење на Коментарот на ЗКП и базата","FINAL MASTER v2.0 · Коментар на ЗКП v6 · JSONL корпус · AI Advokat · Human Gate")
notice(d,"Коментарот е авторски професионален слој, не извор на право. Пред професионална употреба мора да се проверат важечката верзија на ЗКП, официјалниот извор, роковите, списите и релевантната судска практика.","FFF1EF","Правна граница.")
d.add_heading("1. Канонска хиерархија — како се чита системот",1)
bullets(d,["Закон / важечка нормативна одредба.","Официјален извор и проверка на верзија/датум.","Релевантна судска практика и ЕСЧП стандарди кога се применливи.","Коментар на ЗКП на адвокат Зоран Стојанкиќ — авторско толкување и практична систематизација.","Практични водичи и работни модели.","AI синтеза ограничена со проверените извори.","Human Gate — финална човечка професионална проверка."],True)
d.add_heading("2. Quick Start — како да поставите добро прашање",1)
notice(d,"ФАЗА + СУБЈЕКТ + ШТО СЕ СЛУЧИЛО + ШТО САКАТЕ ДА ПОСТИГНЕТЕ","EAF5EF","Формула:")
d.add_paragraph("Наместо „Притвор?“: „Во истрага сум; обвинетиот е во притвор. Кои одредби да ги проверам за преиспитување или замена на мерката?“")
d.add_paragraph("Наместо „Незаконит доказ“: „На главна расправа се предлага доказ од претрес за кој одбраната тврди дека е незаконит. Кои членови и приговори треба да се проверат?“")
p=d.add_heading("3. Што содржат Коментарот и базата",1); p.paragraph_format.page_break_before=True
table(d,["Слој","Содржина","Улога"],[
["Законски текст","Интегрален текст на членот","Нормативна основа во рамките на внесената верзија"],
["Коментар","Правна смисла и објаснување","Авторски секундарен професионален слој"],
["Практика","Примена, стратегија и процесни ризици","Работна ориентација за професионалец"],
["Формулации","Модели за приговор, барање, жалба или аргументација","Почетна рамка што мора да се приспособи"],
["Метаподатоци","Фаза, актери, теми, ризик и поврзани членови","Пребарување, рангирање и насочување"]])
d.add_heading("4. Седумчекорен одговор",1)
table(d,["Чекор","Контролно прашање","Излез"],[
["1","Во која процесна фаза е предметот?","Фаза и контекст"],["2","Кој е релевантниот субјект?","Обвинет, бранител, ЈО, суд, оштетен..."],["3","Кој член е директно релевантен?","Главен член + тесен сет поврзани одредби"],["4","Што пишува законскиот текст?","Нормативен слој"],["5","Што објаснува Коментарот?","Авторски коментар и правна смисла"],["6","Што практично се проверува/прави?","Стратегија, ризици и работна формулација"],["7","Што мора човек да потврди?","Верзија, рок, факти, докази, практика и спис"]])
d.add_heading("5. Четири начини на работа",1)
bullets(d,["Навигација — кога не е познат членот: фаза + актер + тема → релевантни членови.","Проверка по член — законски текст + коментар + поврзани одредби.","Стратегија — практична примена + грешки + ризици + работна формулација.","Редакција / QA — контрола на метаподатоци, упатувања, provenance, верзионирање и release status."])
d.add_heading("6. Практични примери и ризик",1)
d.add_paragraph("Член 2 — презумпција на невиност: одговорот прво го покажува законскиот слој, а потоа коментарот. Human Gate ја проверува точната изјава, контекстот и процесниот статус.")
d.add_paragraph("Член 12 — незаконит доказ: AI прво ја утврдува доказната линија, начинот на прибавување, процесната фаза и поврзаните одредби, па потоа нуди работна рамка за приговор или издвојување.")
d.add_paragraph("Член 329 — признавање вина: нема автоматска препорака „признај/не признавај“. Се анализираат доказите, квалификацијата, последиците и роковите, со задолжителна човечка проценка.")
d.add_paragraph("Fail-closed: ако верзијата, датумот, членот или изворот не може да се потврди, системот не го претставува заклучокот како сигурно тековно право.")
d.add_heading("7. Техничка картичка — JSONL v6",1)
table(d,["Поле","Вредност"],[["Records / unique articles","570 / 570"],["Coverage","1–568 + 267-а + 567-а"],["law_text / commentary","570 / 570"],["Enriched metadata","542"],["Missing enrichment","1–26, 267-а, 567-а"],["Risk distribution","very_high 215 · high 235 · medium 92"],["Canonical role","knowledge_role = authorial_analysis · primary_law_dependency = mk:zkp · brand = AI Advokat"]])
d.add_paragraph("Празно поле не значи дозвола за дополнување по претпоставка. AI мора да работи со фактички достапната структура и да означи што недостасува.")
d.add_heading("8. Human Gate",1)
d.add_paragraph("Човечката проверка е суштинска особено кај лишување од слобода, признавање вина, правни лекови и рокови, незаконити докази, претреси и други дејствија што директно ја менуваат процесната положба.")
notice(d,"AI Advokat може да организира, најде, поврзе и подготви работен материјал. Финалната професионална проценка, потпишување, поднесување или процесна одлука остануваат човечки.","EAF5EF","Human Gate не е одбивање.")
d.add_heading("9. Release protocol",1)
bullets(d,["Проверен извор и exact source hash.","Проверена покриеност на сите 570 записи.","Одвоен законски и авторски слој.","Human Gate approval за release.","D1 read-back со точен број записи и unique articles.","Production deploy и live smoke."],True)
d.add_paragraph("Овој документ ја опишува методологијата; самото постоење на водичот не е доказ дека конкретен production release е активен.")
d.add_heading("10. Nenad Lex Nova Academy",1)
d.add_paragraph("Водичот е основа за програмата „Практична кривична постапка со AI Advokat“: најди го членот, разликувај закон од коментар, работи по процесна фаза, поврзи релевантни одредби и препознај кога е задолжителен Human Gate.")
save(d,"VODIC_ZA_KORISTENJE_KOMENTAR_ZKP_I_BAZA_FINAL_MASTER_v2.0.docx")

d=base_doc("ЗКП Quick Start v2.0","Практичен почеток за прашања во AI Advokat")
notice(d,"ФАЗА + СУБЈЕКТ + ШТО СЕ СЛУЧИЛО + ШТО САКАТЕ ДА ПОСТИГНЕТЕ.","EAF5EF","Формула за добро прашање:")
d.add_heading("Што треба да добиете",1)
bullets(d,["Релевантен член/членови.","Законски слој — одвоен од коментарот.","Авторски коментар на Зоран Стојанкиќ.","Практична примена, ризици и поврзани членови.","Работна формулација кога е соодветно.","Јасна Human Gate напомена."],True)
d.add_heading("Канонска хиерархија",1); d.add_paragraph("Закон/официјален извор → судска практика → Коментар на ЗКП → практични водичи → AI синтеза → Human Gate.")
d.add_heading("Важно",1); d.add_paragraph("AI-помошта не ја заменува проверката на важечката верзија на законот, роковите, списите и професионалната проценка.")
save(d,"ZKP_QUICK_START_v2.docx")

d=base_doc("ZKP Schema Notes v2.0","Техничка картичка на 570-записниот корпус",author="AI Advokat")
table(d,["Поле","Вредност"],[["Records","570"],["Unique articles","570"],["Coverage","1-568 + 267-а + 567-а"],["Enriched","542"],["Missing enrichment","1-26, 267-а, 567-а"],["Risk distribution","very_high 215; high 235; medium 92"],["law_text","570"],["commentary","570"],["practice","567"],["formulations","377"],["related_provisions","544"],["related_articles","538"],["knowledge_role / primary_law_dependency / brand","570 / 570 / 570"],["Original source SHA-256","5a6a691cf616599c1648ab9749a34d388f7c4b5fea8b49ecfd3b4f409903c363"],["Brand-clean SHA-256","d933779ef07d3f6a21e1a2e3593d79610707ce8b06439bbbda005ea3bfa3d860"]])
notice(d,"Primary law first; commentary is authorial secondary material; Human Gate required for professional use.","FFF8E9","Authority model:")
save(d,"ZKP_SCHEMA_NOTES_v2.0.docx")

d=base_doc("ZKP Commentary v6 — Portal Integration","Governed secondary professional layer · 19.10.2026",author="AI Advokat")
d.add_heading("Purpose",1); d.add_paragraph("Register the Коментар на Законот за кривичната постапка · Верзија 6 by Attorney Zoran Stojankich as a governed secondary professional layer in AI Advokat.")
d.add_heading("Boundary",1)
bullets(d,["The existing governed ZKP statutory corpus remains the controlling law layer.","The Commentary is authorial analysis, not a primary legal source.","AI retrieval order is: statutory article → attributed commentary → separately verified case law → labelled AI synthesis.","Professional reliance remains subject to Human Gate.","Canonical brand is AI Advokat / АИ АДВОКАТ. The generation suffix “4” is removed by the normalisation script."],True)
d.add_heading("Delivered source identity",1)
table(d,["Item","Value"],[["DOCX SHA-256","834e9d6fea81cd3c8b3f4fb30bb76186dc199680f60b464a3b9a82457560a3dd"],["JSONL SHA-256","5a6a691cf616599c1648ab9749a34d388f7c4b5fea8b49ecfd3b4f409903c363"],["JSONL records","570"]])
p=d.add_heading("Portal state in this change",1); p.paragraph_format.page_break_before=True
bullets(d,["Registered in data/professional-library.json.","Registered in publication governance and the professional legal corpus.","Added a machine-readable manifest.","Added central-prompt source-separation rules.","Production RAG write remains fail-closed until a separate import Human Gate and read-back verification."])
save(d,"ZKP_COMMENTARY_V6_PORTAL_INTEGRATION_2026-10-19.docx")
