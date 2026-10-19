#!/usr/bin/env python3
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"downloads"/"zkp-commentary-usage"
OUT.mkdir(parents=True, exist_ok=True)

GUIDE_TITLE="ВОДИЧ ЗА КОРИСТЕЊЕ НА КОМЕНТАРОТ НА ЗКП И БАЗАТА"
SUBTITLE="Практичен работен модел за Коментарот v6, базата од 570 записи и AI Advokat"
QUICK_TITLE="КОМЕНТАР НА ЗКП + БАЗА · БРЗА УПОТРЕБА"
AUTHOR="Адвокат Зоран Стојанкиќ · AI Advokat · Скопје · 2026"
SLOGAN="AI Advokat · Правото, организирано. AI, под човечка контрола."

FONT="/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_B="/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
pdfmetrics.registerFont(TTFont("DejaVu", FONT))
pdfmetrics.registerFont(TTFont("DejaVu-Bold", FONT_B))

NAVY=colors.HexColor("#102A46"); GOLD=colors.HexColor("#B9985B"); PALE=colors.HexColor("#EEF3F8"); INK=colors.HexColor("#102033"); MUTED=colors.HexColor("#5A6878"); NOTE=colors.HexColor("#FFF8E8")

seven=[
"Определи ја процесната фаза. Не почнувај само од клучен збор; утврди каде се наоѓа предметот.",
"Пронајди го релевантниот член. Провери го законскиот текст и неговиот статус/верзија.",
"Прочитај го авторскиот коментар. Одделно од законскиот текст: правна смисла, практична примена и стратегија.",
"Провери поврзани одредби. Еден член ретко функционира изолирано.",
"Провери судска практика и официјални извори. Особено кога се потпираш на стандард, рок или спорно толкување.",
"Користи AI Advokat за синтеза. Барај јасно раздвојување: закон → коментар → практика → заклучок.",
"Human Gate. Пред поднесок, совет или процесна одлука, адвокатот ја проверува применливоста врз конкретните факти и доказите."
]
four=[
("По член","Кога е познат релевантниот член од ЗКП."),
("По процесна фаза","Предистражна постапка, истрага, обвинение, главна расправа, пресуда, правни лекови или посебни постапки."),
("По практичен проблем","Притвор, претрес, незаконит доказ, сведок, вештачење, жалба и слично."),
("По адвокатска задача","Приговор, предлог, барање, жалба, стратегија или работна формулација.")
]
schema=["article_number","article_title","law_text","commentary","practice","common_errors","formulations","related_provisions","systematics","procedure_phase","actors","topics","related_articles","practical_use","risk_level"]
checklist=[
"Процесната фаза е правилно определена.",
"Релевантниот член е идентификуван.",
"Законскиот текст е проверен во официјален и важечки извор.",
"Поврзаните одредби се разгледани.",
"Релевантната судска практика е проверена и има provenance.",
"Авторскиот коментар е јасно одделен од законскиот текст.",
"AI синтезата не воведува непотврдени факти или извори.",
"Работната формулација е приспособена на конкретниот предмет.",
"Human Gate е завршен пред надворешна употреба."
]

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcPr.append(shd)

def set_cell_text(cell,text,bold=False,color=None,size=9,center=False):
    cell.text=""
    p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    r=p.add_run(text); r.bold=bold; r.font.name="DejaVu Sans"; r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def doc_base(title):
    d=Document()
    sec=d.sections[0]; sec.top_margin=Inches(.6); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
    normal=d.styles["Normal"]; normal.font.name="DejaVu Sans"; normal.font.size=Pt(10)
    for s in ["Title","Heading 1","Heading 2","Heading 3"]:
        d.styles[s].font.name="DejaVu Sans"; d.styles[s].font.color.rgb=RGBColor(16,42,70)
    header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rr=header.add_run("AI Advokat · Правна едиција"); rr.bold=True; rr.font.name="DejaVu Sans"; rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(90,104,120)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rf=footer.add_run(title); rf.font.name="DejaVu Sans"; rf.font.size=Pt(8); rf.font.color.rgb=RGBColor(90,104,120)
    return d

def add_h(d,text,level=1):
    p=d.add_paragraph(style=f"Heading {level}"); p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(6)
    r=p.add_run(text); r.bold=True
    return p

def add_bullet(d,boldlead,text):
    p=d.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(boldlead); r.bold=True; p.add_run(": "+text)

def make_guide_docx():
    d=doc_base("Водич за користење на Коментарот на ЗКП и базата")
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(35)
    r=p.add_run("AI ADVOKAT"); r.bold=True; r.font.name="DejaVu Sans"; r.font.size=Pt(16); r.font.color.rgb=RGBColor(185,152,91)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(GUIDE_TITLE); r.bold=True; r.font.name="DejaVu Sans"; r.font.size=Pt(25); r.font.color.rgb=RGBColor(16,42,70)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(SUBTITLE); r.font.name="DejaVu Sans"; r.font.size=Pt(13); r.font.color.rgb=RGBColor(90,104,120)
    t=d.add_table(rows=2,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    vals=[("Верзија","2.0 · 2026"),("Опфат","570 записи"),("Законски опфат","чл. 1–568 + 267-а + 567-а"),("Контрола","Human Gate")]
    for i,(a,b) in enumerate(vals):
        c=t.cell(i//2,i%2); shade(c,"EEF3F8"); set_cell_text(c,a+"\n"+b,bold=True,color="173D66",size=10,center=True)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(25)
    r=p.add_run(AUTHOR); r.font.name="DejaVu Sans"; r.font.size=Pt(10); r.font.color.rgb=RGBColor(90,104,120)

    d.add_page_break()
    p=d.add_paragraph(); shade_like=False
    r=p.add_run("ОСНОВНО ПРАВИЛО\n"); r.bold=True
    p.add_run("Коментарот и базата се секундарен авторски слој. Важечкиот закон и официјалните/судските извори остануваат контролен авторитет. AI одговорот е работна синтеза и не ја заменува професионалната проверка.")
    add_h(d,"1. Хиерархија на користење")
    t=d.add_table(rows=1,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(["1 · Закон и официјален извор","2 · Судска практика","3 · Коментар на ЗКП","4 · AI синтеза + Human Gate"]):
        shade(t.cell(0,i),"EEF3F8"); set_cell_text(t.cell(0,i),x,bold=True,size=8,center=True)
    d.add_paragraph("Кога постои разлика меѓу авторскиот коментар и официјалниот важечки текст, предност има официјалниот извор.")
    add_h(d,"2. Четири режими на работа")
    for a,b in four: add_bullet(d,a,b)
    add_h(d,"3. Што содржи еден запис")
    t=d.add_table(rows=5,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(schema):
        c=t.cell(i//3,i%3); shade(c,"102A46"); set_cell_text(c,x,color="FFFFFF",size=8,center=True)
    d.add_paragraph("Кај обогатените записи дополнителните метаподатоци овозможуваат пребарување по фаза, актери, теми, поврзани членови и ризик.")

    d.add_page_break()
    add_h(d,"4. Седумчекорен работен тек")
    t=d.add_table(rows=7,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(seven,1):
        shade(t.cell(i-1,0),"102A46"); set_cell_text(t.cell(i-1,0),str(i),bold=True,color="FFFFFF",size=11,center=True)
        shade(t.cell(i-1,1),"EEF3F8"); set_cell_text(t.cell(i-1,1),x,bold=False,size=9)

    d.add_page_break()
    add_h(d,"5. Quick Start")
    examples=[
      ("Пример 1 — познат член","„Прикажи го член 12 од ЗКП. Оддели го законскиот текст од Коментарот, поврзаните членови и практичната адвокатска примена.“"),
      ("Пример 2 — практичен проблем","„Кои одредби и процесни ризици се релевантни за издвојување незаконит доказ? Наведи го официјалниот законски слој прво.“"),
      ("Пример 3 — правен лек","„Подготви работна анализа за жалбен основ, но означи кои делови доаѓаат од ЗКП, кои од Коментарот и кои се AI синтеза.“")
    ]
    for a,b in examples:
        p=d.add_paragraph(); rr=p.add_run(a+"\n"); rr.bold=True; p.add_run(b)
    add_h(d,"6. Што не смее да се прави")
    for x in ["Не се цитира AI одговор како да е закон.","Не се претставува авторскиот коментар како официјален извор.","Не се користи работна формулација без приспособување на фактите и доказите.","Не се претпоставува дека верзијата на законскиот текст е тековна без проверка на важечката состојба.","Не се меша судска практика со авторска интерпретација без provenance."]:
        d.add_paragraph(x,style="List Bullet")
    add_h(d,"7. Статус на базата и production користење")
    d.add_paragraph("Опфат: 570 уникатни записи - членовите 1–568, плус 267-а и 567-а. Секој запис има законски текст и авторски коментар. Коментарот е регистриран како secondary authorial knowledge layer.")
    d.add_paragraph("Production AI користењето е fail-closed: Коментарот не смее да се вклучи како активен runtime слој додека production ingest, read-back проверката и Human Gate активацијата не завршат. До тогаш овој водич ја дефинира правилната методологија на употреба, но не тврди дека секој запис е веќе live во ботот.")

    d.add_page_break()
    add_h(d,"8. Препорачана форма на AI одговор")
    t=d.add_table(rows=4,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(["A. Релевантен законски текст","B. Верзија / извор","C. Авторски коментар","D. Поврзани одредби","E. Судска практика","F. Практична примена","G. Ризици / отворени прашања","H. Human Gate"]):
        shade(t.cell(i//2,i%2),"EEF3F8"); set_cell_text(t.cell(i//2,i%2),x,bold=True,size=9)
    d.add_paragraph("Овој редослед е дизајниран за да не се изгуби границата меѓу примарното право, секундарната анализа и AI синтезата.")
    add_h(d,"9. Контролна листа пред професионална употреба")
    for x in checklist: d.add_paragraph("☐ "+x)
    p=d.add_paragraph(); rr=p.add_run("ПРАВНА НАПОМЕНА\n"); rr.bold=True
    p.add_run("Овој водич е методолошки и професионален помошен материјал. Не е официјален извор на право и не заменува проверка на важечкиот закон, судската практика и конкретната процесна состојба. Финалната професионална одлука ја носи човек - адвокатот или другото овластено стручно лице.")
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=p.add_run(SLOGAN); rr.bold=True
    d.save(OUT/"VODIC_ZA_KORISTENJE_NA_KOMENTAROT_ZKP_I_BAZATA_v2_FINAL.docx")

def make_quick_docx():
    d=doc_base("AI Advokat · ZKP Quick Start")
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("AI ADVOKAT · ZKP QUICK START"); r.bold=True; r.font.name="DejaVu Sans"; r.font.size=Pt(11); r.font.color.rgb=RGBColor(185,152,91)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(QUICK_TITLE); r.bold=True; r.font.size=Pt(20); r.font.color.rgb=RGBColor(16,42,70)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("570 записи · чл. 1–568 + 267-а + 567-а · Human Gate").bold=True
    add_h(d,"1. Редослед")
    t=d.add_table(rows=1,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(["Закон / официјален извор","Судска практика","Коментар на ЗКП","AI синтеза + Human Gate"]):
        shade(t.cell(0,i),"EEF3F8"); set_cell_text(t.cell(0,i),x,bold=True,size=8,center=True)
    add_h(d,"2. Седум чекори")
    for i,x in enumerate(seven,1): d.add_paragraph(f"{i}. {x.split('.')[0]}.")
    add_h(d,"3. Пример промпт")
    d.add_paragraph("„Прикажи го член 12 од ЗКП. Оддели: (1) законски текст, (2) верзија/извор, (3) Коментар, (4) поврзани членови, (5) судска практика, (6) практична примена и (7) Human Gate.“")
    add_h(d,"4. Не прави")
    for x in ["Не цитирај AI одговор како закон.","Не претставувај авторски коментар како официјален извор.","Не користи работна формулација без приспособување на конкретниот предмет.","Не претпоставувај дека законскиот текст е тековен без официјална проверка."]: d.add_paragraph(x,style="List Bullet")
    p=d.add_paragraph(); rr=p.add_run("КОНТРОЛА: "); rr.bold=True; p.add_run("Коментарот и базата се секундарен авторски слој. Важечкиот закон и официјалните/судските извори се контролен авторитет. Финалната професионална одлука ја носи човек.")
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=p.add_run(SLOGAN); rr.bold=True
    d.save(OUT/"ZKP_QUICK_START_v2.docx")

def pdf_styles():
    return {
      "body":ParagraphStyle("body",fontName="DejaVu",fontSize=9.5,leading=13,textColor=INK,spaceAfter=5),
      "small":ParagraphStyle("small",fontName="DejaVu",fontSize=8,leading=10,textColor=MUTED),
      "h1":ParagraphStyle("h1",fontName="DejaVu-Bold",fontSize=18,leading=22,textColor=NAVY,spaceBefore=8,spaceAfter=8),
      "h2":ParagraphStyle("h2",fontName="DejaVu-Bold",fontSize=14,leading=18,textColor=NAVY,spaceBefore=8,spaceAfter=6),
      "cover":ParagraphStyle("cover",fontName="DejaVu-Bold",fontSize=24,leading=29,textColor=NAVY,alignment=TA_CENTER,spaceAfter=12),
      "sub":ParagraphStyle("sub",fontName="DejaVu",fontSize=12,leading=16,textColor=MUTED,alignment=TA_CENTER,spaceAfter=16),
      "gold":ParagraphStyle("gold",fontName="DejaVu-Bold",fontSize=13,leading=16,textColor=GOLD,alignment=TA_CENTER,spaceAfter=12),
      "center":ParagraphStyle("center",fontName="DejaVu",fontSize=9,leading=12,textColor=MUTED,alignment=TA_CENTER),
      "bold":ParagraphStyle("bold",fontName="DejaVu-Bold",fontSize=9.3,leading=12,textColor=INK),
    }

def footer(canvas,doc):
    canvas.saveState(); canvas.setFont("DejaVu",7.5); canvas.setFillColor(MUTED)
    canvas.drawRightString(A4[0]-18*mm,A4[1]-10*mm,"AI Advokat · Правна едиција")
    canvas.drawCentredString(A4[0]/2,10*mm,f"Водич за користење на Коментарот на ЗКП и базата · {doc.page}")
    canvas.restoreState()

def make_guide_pdf():
    st=pdf_styles(); path=OUT/"VODIC_ZA_KORISTENJE_NA_KOMENTAROT_ZKP_I_BAZATA_v2_FINAL.pdf"
    doc=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=14*mm,leftMargin=14*mm,topMargin=18*mm,bottomMargin=18*mm)
    S=[]
    S += [Spacer(1,18*mm),Paragraph("AI ADVOKAT",st["gold"]),Paragraph(GUIDE_TITLE,st["cover"]),Paragraph(SUBTITLE,st["sub"])]
    data=[[Paragraph("<b>Верзија</b><br/>2.0 · 2026",st["center"]),Paragraph("<b>Опфат</b><br/>570 записи",st["center"])],[Paragraph("<b>Законски опфат</b><br/>чл. 1–568 + 267-а + 567-а",st["center"]),Paragraph("<b>Контрола</b><br/>Human Gate",st["center"])]]
    t=Table(data,colWidths=[82*mm,82*mm],rowHeights=[23*mm,23*mm]); t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),PALE),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("BOX",(0,0),(-1,-1),.4,colors.white)])); S += [t,Spacer(1,20*mm),Paragraph(AUTHOR,st["center"]),PageBreak()]
    note=Table([[Paragraph("<b>ОСНОВНО ПРАВИЛО</b><br/>Коментарот и базата се секундарен авторски слој. Важечкиот закон и официјалните/судските извори остануваат контролен авторитет. AI одговорот е работна синтеза и не ја заменува професионалната проверка.",st["body"])]],colWidths=[174*mm]); note.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),NOTE),("BOX",(0,0),(-1,-1),.8,GOLD),("LEFTPADDING",(0,0),(-1,-1),8),("RIGHTPADDING",(0,0),(-1,-1),8),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7)])); S += [note,Spacer(1,6*mm),Paragraph("1. Хиерархија на користење",st["h1"])]
    arr=[[Paragraph(x,st["bold"]) for x in ["1 · Закон и официјален извор","2 · Судска практика","3 · Коментар на ЗКП","4 · AI синтеза + Human Gate"]]]
    t=Table(arr,colWidths=[43.5*mm]*4); t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),PALE),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ALIGN",(0,0),(-1,-1),"CENTER"),("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4)])); S += [t,Paragraph("Кога постои разлика меѓу авторскиот коментар и официјалниот важечки текст, предност има официјалниот извор.",st["body"]),Paragraph("2. Четири режими на работа",st["h1"])]
    for a,b in four: S.append(Paragraph("• <b>"+a+":</b> "+b,st["body"]))
    S += [Paragraph("3. Што содржи еден запис",st["h1"])]
    rows=[]; 
    for i in range(0,len(schema),3): rows.append([Paragraph(x,ParagraphStyle("sc",fontName="DejaVu",fontSize=7.7,leading=9,textColor=colors.white,alignment=TA_CENTER)) for x in schema[i:i+3]])
    t=Table(rows,colWidths=[58*mm]*3,rowHeights=10*mm); t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),NAVY),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("BOX",(0,0),(-1,-1),.3,NAVY)])); S += [t,Paragraph("Кај обогатените записи дополнителните метаподатоци овозможуваат пребарување по фаза, актери, теми, поврзани членови и ризик.",st["body"]),PageBreak()]
    S += [Paragraph("4. Седумчекорен работен тек",st["h1"])]
    rows=[]
    for i,x in enumerate(seven,1):
        rows.append([Paragraph(f"<b>{i}</b>",ParagraphStyle("n",fontName="DejaVu-Bold",fontSize=10,textColor=colors.white,alignment=TA_CENTER)),Paragraph(x,st["body"])])
    t=Table(rows,colWidths=[15*mm,159*mm]); t.setStyle(TableStyle([("BACKGROUND",(0,0),(0,-1),NAVY),("BACKGROUND",(1,0),(1,-1),PALE),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7),("GRID",(0,0),(-1,-1),2,colors.white)])); S += [t,PageBreak()]
    S += [Paragraph("5. Quick Start",st["h1"])]
    ex=[("Пример 1 — познат член","„Прикажи го член 12 од ЗКП. Оддели го законскиот текст од Коментарот, поврзаните членови и практичната адвокатска примена.“"),("Пример 2 — практичен проблем","„Кои одредби и процесни ризици се релевантни за издвојување незаконит доказ? Наведи го официјалниот законски слој прво.“"),("Пример 3 — правен лек","„Подготви работна анализа за жалбен основ, но означи кои делови доаѓаат од ЗКП, кои од Коментарот и кои се AI синтеза.“")]
    for a,b in ex: S += [Paragraph("<b>"+a+"</b>",st["body"]),Paragraph(b,st["body"]),Spacer(1,2*mm)]
    S += [Paragraph("6. Што не смее да се прави",st["h1"])]
    for x in ["Не се цитира AI одговор како да е закон.","Не се претставува авторскиот коментар како официјален извор.","Не се користи работна формулација без приспособување на фактите и доказите.","Не се претпоставува дека верзијата на законскиот текст е тековна без проверка на важечката состојба.","Не се меша судска практика со авторска интерпретација без provenance."]: S.append(Paragraph("• "+x,st["body"]))
    S += [Paragraph("7. Статус на базата и production користење",st["h1"]),Paragraph("Опфат: 570 уникатни записи - членовите 1–568, плус 267-а и 567-а. Секој запис има законски текст и авторски коментар. Коментарот е регистриран како secondary authorial knowledge layer.",st["body"]),Paragraph("Production AI користењето е fail-closed: Коментарот не смее да се вклучи како активен runtime слој додека production ingest, read-back проверката и Human Gate активацијата не завршат. До тогаш овој водич ја дефинира правилната методологија на употреба, но не тврди дека секој запис е веќе live во ботот.",st["body"]),PageBreak()]
    S += [Paragraph("8. Препорачана форма на AI одговор",st["h1"])]
    vals=["A. Релевантен законски текст","B. Верзија / извор","C. Авторски коментар","D. Поврзани одредби","E. Судска практика","F. Практична примена","G. Ризици / отворени прашања","H. Human Gate"]
    rows=[[Paragraph("<b>"+vals[i]+"</b>",st["body"]),Paragraph("<b>"+vals[i+1]+"</b>",st["body"])] for i in range(0,8,2)]
    t=Table(rows,colWidths=[87*mm,87*mm]); t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),PALE),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7)])); S += [t,Paragraph("Овој редослед е дизајниран за да не се изгуби границата меѓу примарното право, секундарната анализа и AI синтезата.",st["body"]),Paragraph("9. Контролна листа пред професионална употреба",st["h1"])]
    for x in checklist: S.append(Paragraph("☐ "+x,st["body"]))
    note=Table([[Paragraph("<b>ПРАВНА НАПОМЕНА</b><br/>Овој водич е методолошки и професионален помошен материјал. Не е официјален извор на право и не заменува проверка на важечкиот закон, судската практика и конкретната процесна состојба. Финалната професионална одлука ја носи човек - адвокатот или другото овластено стручно лице.",st["body"])]],colWidths=[174*mm]); note.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),NOTE),("BOX",(0,0),(-1,-1),.8,GOLD),("LEFTPADDING",(0,0),(-1,-1),8),("RIGHTPADDING",(0,0),(-1,-1),8),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7)])); S += [Spacer(1,3*mm),note,Spacer(1,6*mm),Paragraph("<b>"+SLOGAN+"</b>",st["center"])]
    doc.build(S,onFirstPage=footer,onLaterPages=footer)

def make_quick_pdf():
    st=pdf_styles(); path=OUT/"ZKP_QUICK_START_v2.pdf"
    doc=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=14*mm,leftMargin=14*mm,topMargin=14*mm,bottomMargin=14*mm)
    q=ParagraphStyle("q",fontName="DejaVu",fontSize=8.5,leading=11,textColor=INK)
    S=[Paragraph("AI ADVOKAT · ZKP QUICK START",st["gold"]),Paragraph(QUICK_TITLE,st["cover"]),Paragraph("<b>570 записи · чл. 1–568 + 267-а + 567-а · Human Gate</b>",st["center"]),Spacer(1,4*mm),Paragraph("1. Редослед",st["h2"])]
    arr=[[Paragraph(x,q) for x in ["Закон / официјален извор","Судска практика","Коментар на ЗКП","AI синтеза + Human Gate"]]]
    t=Table(arr,colWidths=[43.5*mm]*4); t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),PALE),("ALIGN",(0,0),(-1,-1),"CENTER"),("VALIGN",(0,0),(-1,-1),"MIDDLE")])); S += [t,Spacer(1,3*mm),Paragraph("2. Седум чекори",st["h2"])]
    for i,x in enumerate(seven,1): S.append(Paragraph(f"{i}. {x.split('.')[0]}.",q))
    S += [Paragraph("3. Пример промпт",st["h2"]),Paragraph("„Прикажи го член 12 од ЗКП. Оддели: (1) законски текст, (2) верзија/извор, (3) Коментар, (4) поврзани членови, (5) судска практика, (6) практична примена и (7) Human Gate.“",q),Paragraph("4. Не прави",st["h2"])]
    for x in ["Не цитирај AI одговор како закон.","Не претставувај авторски коментар како официјален извор.","Не користи работна формулација без приспособување на конкретниот предмет.","Не претпоставувај дека законскиот текст е тековен без официјална проверка."]: S.append(Paragraph("• "+x,q))
    note=Table([[Paragraph("<b>КОНТРОЛА:</b> Коментарот и базата се секундарен авторски слој. Важечкиот закон и официјалните/судските извори се контролен авторитет. Финалната професионална одлука ја носи човек.",q)]],colWidths=[174*mm]); note.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),NOTE),("BOX",(0,0),(-1,-1),.8,GOLD),("LEFTPADDING",(0,0),(-1,-1),7),("RIGHTPADDING",(0,0),(-1,-1),7),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)])); S += [Spacer(1,3*mm),note,Spacer(1,4*mm),Paragraph("<b>"+SLOGAN+"</b>",st["center"])]
    doc.build(S)

if __name__=="__main__":
    make_guide_docx(); make_quick_docx(); make_guide_pdf(); make_quick_pdf()
    for p in sorted(OUT.iterdir()):
        if p.is_file(): print(p.name,p.stat().st_size)
